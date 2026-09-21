"""Unit tests for DOCX document parser."""

import io
import uuid
import docx
import pytest
from app.pipeline.docx import DOCXDocumentParser
from app.pipeline.schemas import BlockType


def make_sample_docx() -> bytes:
    """Generate a sample DOCX document in memory with headings, text, and tables."""
    doc = docx.Document()
    doc.add_heading("CRPF Procurement Guidelines", level=1)
    doc.add_paragraph("This document outlines requirements for procurement of security equipment.")
    
    table = doc.add_table(rows=2, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Item Code"
    hdr_cells[1].text = "Description"
    hdr_cells[2].text = "Qty"

    row_cells = table.rows[1].cells
    row_cells[0].text = "CRPF-001"
    row_cells[1].text = "Tactical Helmet"
    row_cells[2].text = "1000"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_parser_can_parse():
    parser = DOCXDocumentParser()
    assert parser.can_parse(".docx") is True
    assert parser.can_parse(".doc") is True
    assert parser.can_parse(".pdf") is False


def test_docx_parser_success():
    parser = DOCXDocumentParser()
    doc_id = uuid.uuid4()
    content = make_sample_docx()

    result = parser.parse(document_id=doc_id, content=content, filename="guidelines.docx")

    assert result.document_id == doc_id
    assert result.document_type == "WORD_DOCUMENT"
    assert result.page_count == 1
    assert result.total_tables == 1
    assert len(result.pages) == 1
    
    blocks = result.pages[0].blocks
    assert len(blocks) >= 3

    heading_blocks = [b for b in blocks if b.type == BlockType.HEADING]
    assert len(heading_blocks) >= 1
    assert "CRPF Procurement Guidelines" in heading_blocks[0].text

    table_blocks = [b for b in blocks if b.type == BlockType.TABLE]
    assert len(table_blocks) == 1
    assert table_blocks[0].table_data is not None
    assert "Item Code" in table_blocks[0].table_data.headers
    assert len(table_blocks[0].table_data.rows) == 1


def test_docx_parser_empty_content_raises():
    parser = DOCXDocumentParser()
    with pytest.raises(ValueError) as exc_info:
        parser.parse(document_id=uuid.uuid4(), content=b"", filename="empty.docx")
    assert "Empty DOCX" in str(exc_info.value)


def test_docx_parser_corrupted_raises():
    parser = DOCXDocumentParser()
    with pytest.raises(ValueError) as exc_info:
        parser.parse(document_id=uuid.uuid4(), content=b"PK\x03\x04corrupted zip payload", filename="bad.docx")
    assert "Corrupted or invalid DOCX" in str(exc_info.value)
