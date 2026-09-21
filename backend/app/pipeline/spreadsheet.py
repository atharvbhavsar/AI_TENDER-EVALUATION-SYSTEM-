"""Spreadsheet parser for extracting structured tables and grid data from XLSX files."""

import io
import logging
import uuid
from typing import List
import openpyxl

from app.core.config import get_settings
from app.pipeline.base import BaseDocumentParser
from app.pipeline.schemas import BlockType, DocumentBlock, DocumentPage, NormalizedDocument, TableData

logger = logging.getLogger("app.pipeline.spreadsheet")


class SpreadsheetDocumentParser(BaseDocumentParser):
    """Parser for Microsoft Excel (.xlsx) spreadsheet workbooks."""

    def can_parse(self, file_extension: str) -> bool:
        return file_extension.lower() == ".xlsx"

    def parse(
        self,
        document_id: uuid.UUID,
        content: bytes,
        filename: str,
    ) -> NormalizedDocument:
        settings = get_settings()

        if not content or len(content) == 0:
            raise ValueError("Empty spreadsheet content.")

        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        except Exception as exc:
            logger.error("Failed to parse Excel workbook '%s': %s", filename, str(exc))
            raise ValueError(f"Corrupted or invalid spreadsheet: {str(exc)}")

        blocks: List[DocumentBlock] = []
        total_chars = 0
        total_tables = 0
        block_idx = 1

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            sheet = wb[sheet_name]

            # 1. Add Sheet Heading block
            blocks.append(
                DocumentBlock(
                    block_id=f"p1_b{block_idx}",
                    type=BlockType.HEADING,
                    text=f"Sheet: {sheet_name}",
                    confidence=1.0,
                )
            )
            block_idx += 1
            total_chars += len(sheet_name)

            # 2. Extract Tabular Grid Data
            grid_rows: List[List[str]] = []
            for row in sheet.iter_rows(values_only=True):
                # Filter out completely empty rows
                row_values = [str(val).strip() if val is not None else "" for val in row]
                if any(row_values):
                    grid_rows.append(row_values)

            if grid_rows:
                headers = grid_rows[0]
                data_rows = grid_rows[1:] if len(grid_rows) > 1 else []
                table_text = " | ".join(headers) + "\n" + "\n".join([" | ".join(r) for r in data_rows])

                blocks.append(
                    DocumentBlock(
                        block_id=f"p1_table_{sheet_idx+1}",
                        type=BlockType.TABLE,
                        text=table_text,
                        confidence=1.0,
                        table_data=TableData(headers=headers, rows=data_rows),
                    )
                )
                total_tables += 1
                total_chars += len(table_text)

        wb.close()

        pages = [
            DocumentPage(
                page_number=1,
                blocks=blocks,
            )
        ]

        return NormalizedDocument(
            document_id=document_id,
            document_type="SPREADSHEET",
            processor_version=settings.PROCESSOR_VERSION,
            page_count=1,
            total_characters=total_chars,
            total_tables=total_tables,
            pages=pages,
            metadata={"filename": filename, "sheets_count": len(wb.sheetnames)},
        )
