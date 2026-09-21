"""DOCX Office document parser for extracting structured paragraphs, headings, and tables."""

import io
import logging
import uuid
from typing import List
import docx

from app.core.config import get_settings
from app.pipeline.base import BaseDocumentParser
from app.pipeline.schemas import BlockType, DocumentBlock, DocumentPage, NormalizedDocument, TableData

logger = logging.getLogger("app.pipeline.docx")


class DOCXDocumentParser(BaseDocumentParser):
    """Parser for Microsoft Word (.docx) documents."""

    def can_parse(self, file_extension: str) -> bool:
        return file_extension.lower() in (".docx", ".doc")

    def parse(
        self,
        document_id: uuid.UUID,
        content: bytes,
        filename: str,
    ) -> NormalizedDocument:
        settings = get_settings()

        if not content or len(content) == 0:
            raise ValueError("Empty DOCX document content.")

        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception as exc:
            logger.error("Failed to parse DOCX document '%s': %s", filename, str(exc))
            raise ValueError(f"Corrupted or invalid DOCX document: {str(exc)}")

        blocks: List[DocumentBlock] = []
        total_chars = 0
        total_tables = 0
        block_idx = 1

        # 1. Extract Paragraphs & Headings
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = paragraph.style.name if paragraph.style else ""
            is_heading = style_name.lower().startswith("heading") or style_name.lower().startswith("title")
            block_type = BlockType.HEADING if is_heading else BlockType.TEXT

            blocks.append(
                DocumentBlock(
                    block_id=f"p1_b{block_idx}",
                    type=block_type,
                    text=text,
                    confidence=1.0,
                )
            )
            block_idx += 1
            total_chars += len(text)

        # 2. Extract Tables
        for table_idx, table in enumerate(doc.tables):
            table_rows: List[List[str]] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(row_cells)

            if table_rows:
                headers = table_rows[0]
                body_rows = table_rows[1:] if len(table_rows) > 1 else []
                table_text = " | ".join(headers) + "\n" + "\n".join([" | ".join(r) for r in body_rows])

                blocks.append(
                    DocumentBlock(
                        block_id=f"p1_table_{table_idx+1}",
                        type=BlockType.TABLE,
                        text=table_text,
                        confidence=1.0,
                        table_data=TableData(headers=headers, rows=body_rows),
                    )
                )
                total_tables += 1
                total_chars += len(table_text)

        # DOCX is represented as a single cohesive structural page
        pages = [
            DocumentPage(
                page_number=1,
                blocks=blocks,
            )
        ]

        return NormalizedDocument(
            document_id=document_id,
            document_type="WORD_DOCUMENT",
            processor_version=settings.PROCESSOR_VERSION,
            page_count=1,
            total_characters=total_chars,
            total_tables=total_tables,
            pages=pages,
            metadata={"filename": filename},
        )
